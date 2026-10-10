# -*- coding: utf-8 -*-
"""Construye el modelo en Excel con fórmulas vivas que reproduce engine.py en modo determinístico,
más hojas de resultados fijos (escenarios, simulación, sensibilidad) exportadas desde Python."""
import json, sys
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.comments import Comment
from inputs import SKUS, P, FIXED, SETUP, ANNUAL_CHARGES, LIC_M, RS_PAY_M, FIXED_KIND, SETUP_KIND, ANNUAL_KIND

OUT = sys.argv[1] if len(sys.argv) > 1 else 'modelo.xlsx'
A = json.load(open('analysis.json'))
ST = json.load(open('structures.json'))   # simulación de 20,000 escenarios por estructura


F = 'Arial'
BLUE = Font(name=F, color='0000FF', size=10)
BLACK = Font(name=F, color='000000', size=10)
GREEN = Font(name=F, color='008000', size=10)
BOLD = Font(name=F, bold=True, size=10)
TITLE = Font(name=F, bold=True, size=14)
SUB = Font(name=F, bold=True, size=11, color='1F3864')
NOTE = Font(name=F, italic=True, size=9, color='595959')
YEL = PatternFill('solid', fgColor='FFFF00')
HEAD = PatternFill('solid', fgColor='D9E1F2')
GREY = PatternFill('solid', fgColor='F2F2F2')
thin = Side(style='thin', color='BFBFBF')
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
NUM = '#,##0;(#,##0);"-"'
NUM2 = '#,##0.00;(#,##0.00);"-"'
NUM4 = '0.0000;(0.0000);"-"'
PCT = '0.0%;(0.0%);"-"'
MULT = '0.00"x"'

wb = Workbook()
names = {}


def name(nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)
    names[nm] = ref


def hdr(ws, row, cols, labels, widths=None):
    for c, lab in zip(cols, labels):
        cell = ws.cell(row=row, column=c, value=lab)
        cell.font = BOLD; cell.fill = HEAD; cell.alignment = Alignment(wrap_text=True, vertical='center'); cell.border = BOX


def setw(ws, widths):
    for k, v in widths.items():
        ws.column_dimensions[k].width = v


def put(ws, r, c, v, font=BLACK, fmt=None, fill=None, wrap=False, comment=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font = font
    if fmt: cell.number_format = fmt
    if fill: cell.fill = fill
    if wrap: cell.alignment = Alignment(wrap_text=True, vertical='top')
    if comment: cell.comment = Comment(comment, 'Modelo')
    return cell

# =============================================================== Léeme
ws = wb.active; ws.title = 'Léeme'
setw(ws, {'A': 110})
lines = [
 ('Modelo financiero: distribuidora de medicamentos para compras públicas en Panamá', TITLE),
 ('Versión auditada del 9 de octubre de 2026. Cifras en balboas (B/. = US$). Mes 1 del modelo = enero de 2027.', NOTE),
 ('', None),
 ('Cómo usarlo', SUB),
 ('1. Elija el escenario en Supuestos!B3 (1 = pesimista, 2 = base, 3 = optimista). Todo el modelo se recalcula.', BLACK),
 ('2. Cambie cualquier valor azul de las columnas D, E o F de Supuestos, o la columna «Incluir» de SKUs, para probar otra hipótesis.', BLACK),
 ('3. Supuestos!B4 = 1 aplica las tarifas de la propuesta de SignatureLex en lugar de las de mercado.', BLACK),
 ('4. Resultados vivos: Valoración (VPN, TIR, múltiplo, capital), Anual (estado de resultados e impuestos), Mensual (flujo de caja).', BLACK),
 ('5. Las hojas Escenarios, Simulación y Sensibilidad son resultados fijos del motor en Python; no cambian al editar supuestos.', BLACK),
 ('', None),
 ('Convenciones', SUB),
 ('Texto azul = dato de entrada editable. Negro = fórmula. Verde = vínculo a otra hoja. Relleno amarillo = supuestos clave.', BLACK),
 ('En la columna Fuente: F = dato verificado en fuente abierta; R = referencia indirecta; E = estimación propia pendiente de cotización.', BLACK),
 ('', None),
 ('Lógica del modelo (mes a mes, 60 meses desde la constitución de la S.A.)', SUB),
 ('Ventas descentralizadas = demanda anual x factor de mercado x precio de oferta x probabilidad de registro x participación x rampa.', BLACK),
 ('Ventas centralizadas = demanda x precio x probabilidad de ganar x cantidad realmente comprada (LPU de la CSS; compra centralizada del MINSA).', BLACK),
 ('La demanda del MINSA centralizado usa la subasta de 2026 cuando existe y, si no, la compra conjunta de 2023 multiplicada por p_cc_w.', BLACK),
 ('Solo se oferta un producto en un canal si el margen bruto supera el mínimo (35% descentralizado, 20% centralizado).', BLACK),
 ('Inventario objetivo = meses de cobertura x costo del mes siguiente; compras con anticipo al ordenar y saldo al embarque, menos el crédito de la fábrica.', BLACK),
 ('Cobro = ventas de hace «días de cobro» meses. Fianzas de cumplimiento sobre la cantidad máxima del contrato (30% LPU, 20% MINSA), vigentes 3 meses (LPU) o 6 (MINSA) después de cada contrato; fianza de propuesta de 10% sobre toda la oferta en la LPU (el acto supera B/.70,000: Ley 419, art. 245); el colateral queda inmovilizado.', BLACK),
 ('El primer contrato centralizado exige el registro al ofertar: si no está a tiempo, pasa al ciclo siguiente de 15 meses. La segunda ola entra en el ciclo siguiente para el que tiene registro.', BLACK),
 ('Costos fijos con cargas laborales sobre regente y coordinador e ITBMS no recuperable (7%) sobre servicios; el coordinador pasa a tiempo completo con el primer contrato centralizado.', BLACK),
 ('Aportes de capital = lo necesario para mantener la caja mínima. Dividendos en los meses 24, 36 y 48 si hay utilidades y caja.', BLACK),
 ('Valor terminal (mes 60): caja + 90% del capital de trabajo (conservador). Es el que usan los documentos.', BLACK),
 ('Impuestos: ISR 25% con cuotas estimadas en junio, septiembre y diciembre y saldo en marzo; pérdidas arrastrables 5 años al 20% anual con tope de 50% de la renta; CAIR si ventas > B/.1.5 millones; dividendos 10%; aviso de operación 2% del patrimonio, mínimo B/.100 y tope B/.60,000.', BLACK),
 ('', None),
 ('Limitaciones', SUB),
 ('La fábrica no ha dado precios: el costo parte de una referencia por producto = precio representativo de las compras públicas indias vigentes x 1.4 (ajustado en cinco productos; ver la columna de fuente en Productos), multiplicada por el índice del proveedor (1.25 en la base).', BLACK),
 ('No incluye ventas al sector privado ni productos que exigen equivalencia terapéutica. Los costos fijos y logísticos son supuestos hasta tener cotizaciones.', BLACK),
 ('El capital del caso base es un valor esperado sin fallas; para dimensionar el capital use los percentiles de la hoja Simulación.', BLACK),
]
for i, (t, f) in enumerate(lines, start=1):
    c = ws.cell(row=i, column=1, value=t)
    if f: c.font = f
    c.alignment = Alignment(wrap_text=True)

# =============================================================== Supuestos
ws = wb.create_sheet('Supuestos')
setw(ws, {'A': 13, 'B': 62, 'C': 9, 'D': 12, 'E': 12, 'F': 12, 'G': 12, 'H': 70, 'I': 10})
put(ws, 1, 1, 'Supuestos', TITLE)
put(ws, 3, 1, 'Escenario', BOLD); put(ws, 3, 2, 2, BLUE, fill=YEL, comment='1 = pesimista, 2 = base, 3 = optimista')
put(ws, 3, 3, '=CHOOSE($B$3,"Pesimista","Base","Optimista")', BLACK)
put(ws, 4, 1, 'SignatureLex', BOLD); put(ws, 4, 2, 0, BLUE, fill=YEL, comment='1 = usar tarifas de la propuesta de SignatureLex; 0 = tarifas de mercado')
put(ws, 4, 3, 'Usar tarifas de la propuesta de SignatureLex (1) o de mercado (0)', NOTE)
name('p_sc', "Supuestos!$B$3"); name('p_sl', "Supuestos!$B$4")
r = 6
hdr(ws, r, range(1, 9), ['Clave', 'Parámetro', 'Unidad', 'Pesimista', 'Base', 'Optimista', 'Usado', 'Fuente o base'])
r += 1
key_rows = {}
key_assumptions = {'sup_idx', 'share_dec', 'pwin_cen', 'dso_days', 'reg_m', 'cen_start', 'credit_days', 'adv', 'mkt', 'n2'}
for k, v in P.items():
    lab, unit, pes, base, opt, src = v
    put(ws, r, 1, k, NOTE)
    put(ws, r, 2, lab, BLACK, wrap=True)
    put(ws, r, 3, unit, BLACK)
    fmt = PCT if unit == '%' else (MULT if unit == 'x' else NUM)
    if k in ('sup_idx', 'log_mult', 'mkt', 'pf_dec', 'pf_cen', 'gc_mult', 'rel2'):
        fmt = '0.00'
    for j, val in enumerate((pes, base, opt)):
        put(ws, r, 4 + j, val, BLUE, fmt, fill=YEL if k in key_assumptions else None)
    if k == 'rs_law':
        f = '=IF(p_sl=1,1350,CHOOSE(p_sc,D{0},E{0},F{0}))'.format(r)
    elif k == 'rs_iea':
        f = '=IF(p_sl=1,1500,CHOOSE(p_sc,D{0},E{0},F{0}))'.format(r)
    else:
        f = '=CHOOSE(p_sc,D{0},E{0},F{0})'.format(r)
    put(ws, r, 7, f, BLACK, fmt)
    put(ws, r, 8, src, NOTE, wrap=True)
    name('p_' + k, f"Supuestos!$G${r}")
    key_rows[k] = r
    r += 1

r += 1
put(ws, r, 1, 'Costos fijos mensuales (B/./mes)', SUB); r += 1
hdr(ws, r, range(1, 11), ['#', 'Concepto', 'Inicio', 'Pesimista', 'Base', 'Optimista', 'Usado (con cargas o ITBMS)', 'Fuente o base', 'Mes inicio', 'Tipo'])
r += 1
f_first = r
KLAB = {'staff': 'sueldo + cargas laborales', 'service': 'servicio + ITBMS', 'other': 'otro'}
KF = {'staff': '*(1+p_oncost)', 'service': '*(1+p_itbms_srv)', 'other': ''}
for i, fx in enumerate(FIXED, start=1):
    lab, pes, base, opt, rule, src = fx
    kd = FIXED_KIND[i - 1]
    put(ws, r, 1, i, NOTE); put(ws, r, 2, lab, BLACK, wrap=True)
    put(ws, r, 3, {'m1': 'mes 1', 'lic': 'licencia', 'prep': 'registros - 1', 'cen': '1er contrato centralizado'}[rule], BLACK)
    for j, val in enumerate((pes, base, opt)):
        put(ws, r, 4 + j, val, BLUE, NUM)
    put(ws, r, 7, f'=CHOOSE(p_sc,D{r},E{r},F{r}){KF[kd]}', BLACK, NUM)
    put(ws, r, 8, src, NOTE, wrap=True)
    start = {'m1': '1', 'lic': 'c_lic_m', 'cen': 'c_cen_start'}.get(rule, 'MAX(c_reg_m-1,1)')
    put(ws, r, 9, '=' + start, BLACK, NUM)
    put(ws, r, 10, KLAB[kd], NOTE)
    r += 1
f_last = r - 1
put(ws, r, 2, 'Total costos fijos a plena operación', BOLD); put(ws, r, 7, f'=SUM(G{f_first}:G{f_last})', BOLD, NUM)
name('f_used', f"Supuestos!$G${f_first}:$G${f_last}"); name('f_start', f"Supuestos!$I${f_first}:$I${f_last}")
name('f_total', f"Supuestos!$G${r}")
r += 2
put(ws, r, 1, 'Gastos de arranque, una sola vez (B/.)', SUB); r += 1
hdr(ws, r, range(1, 11), ['#', 'Concepto', 'Mes', 'Pesimista', 'Base', 'Optimista', 'Usado (con ITBMS)', 'Fuente o base', 'SignatureLex', 'Tipo'])
r += 1
u_first = r
for i, s in enumerate(SETUP, start=1):
    lab, pes, base, opt, month, slq, src = s
    kd = SETUP_KIND[i - 1]
    put(ws, r, 1, i, NOTE); put(ws, r, 2, lab, BLACK, wrap=True); put(ws, r, 3, month, BLUE, NUM)
    for j, val in enumerate((pes, base, opt)):
        put(ws, r, 4 + j, val, BLUE, NUM)
    put(ws, r, 9, slq, BLUE, NUM)
    put(ws, r, 7, f'=IF(AND(p_sl=1,I{r}>0),I{r},CHOOSE(p_sc,D{r},E{r},F{r})){KF[kd]}', BLACK, NUM)
    put(ws, r, 8, src, NOTE, wrap=True)
    put(ws, r, 10, KLAB[kd], NOTE)
    r += 1
u_last = r - 1
put(ws, r, 2, 'Total arranque (incluye fase 0)', BOLD); put(ws, r, 7, f'=SUM(G{u_first}:G{u_last})', BOLD, NUM)
name('u_used', f"Supuestos!$G${u_first}:$G${u_last}"); name('u_month', f"Supuestos!$C${u_first}:$C${u_last}")
put(ws, r + 1, 2, 'Fase 0 (filas con mes 0)', BOLD); put(ws, r + 1, 7, '=SUMIF(u_month,0,u_used)', BOLD, NUM)
name('c_phase0', f"Supuestos!$G${r + 1}")
r += 1
r += 2
put(ws, r, 1, 'Cargos anuales (meses 1, 13, 25, 37 y 49; el aviso de operación se calcula aparte sobre el patrimonio)', SUB); r += 1
a_first = r
for (lab, val, src), kd in zip(ANNUAL_CHARGES, ANNUAL_KIND):
    put(ws, r, 2, lab, BLACK); put(ws, r, 5, val, BLUE, NUM); put(ws, r, 7, f'=E{r}{KF[kd]}', BLACK, NUM); put(ws, r, 8, src, NOTE)
    put(ws, r, 10, KLAB[kd], NOTE); r += 1
put(ws, r, 2, 'Total cargos anuales', BOLD); put(ws, r, 7, f'=SUM(G{a_first}:G{r-1})', BOLD, NUM); name('c_annual', f"Supuestos!$G${r}")
r += 2
put(ws, r, 1, 'Valores calculados (no editar)', SUB); r += 1
calc = [
 ('c_lic_m', 'Mes de la licencia de operación', str(LIC_M)),
 ('c_rs_pay_m', 'Mes en que se pagan los registros de la primera ola', str(RS_PAY_M)),
 ('c_lead_m', 'Tiempo de reposición (meses)', 'ROUND(p_lead_days/30,0)'),
 ('c_dso_m', 'Días de cobro (meses)', 'ROUND(p_dso_days/30,0)'),
 ('c_credit_m', 'Crédito del proveedor (meses)', 'ROUND(p_credit_days/30,0)'),
 ('c_ship_m', 'Embarque a stock (meses)', 'ROUND(p_ship_m,0)'),
 ('c_reg_m', 'Mes de registros (entero)', 'ROUND(p_reg_m,0)'),
 ('c_first_dec', 'Primer mes con ventas descentralizadas', 'c_reg_m+c_lead_m'),
 ('c_cen_start', 'Primer mes con contrato centralizado (el registro debe existir al ofertar, c_bid_m meses antes)', 'MAX(ROUND(p_cen_start,0)+15*ROUNDUP(MAX(0,c_reg_m+c_bid_m-ROUND(p_cen_start,0))/15,0),c_first_dec)'),
 ('c_rt', 'Meses de trámite del registro', 'MAX(c_reg_m-c_rs_pay_m,1)'),
 ('c_m2', 'Mes de registros de la segunda ola', 'ROUND(p_m2,0)'),
 ('c_first_dec2', 'Primer mes con ventas de la segunda ola', 'c_m2+c_rt+c_lead_m'),
 ('c_c2_lpu', 'Primer contrato de la LPU con productos de la segunda ola (siguiente ciclo de 15 meses)', 'c_cen_start+15*ROUNDUP(MAX(0,MAX(c_first_dec2,c_m2+c_rt+c_bid_m)-c_cen_start)/15,0)'),
 ('c_c2_minsa', 'Primer contrato del MINSA con productos de la segunda ola (siguiente ciclo de 12 meses)', 'c_cen_start+12*ROUNDUP(MAX(0,MAX(c_first_dec2,c_m2+c_rt+1)-c_cen_start)/12,0)'),
 ('c_n1', 'Productos presentados en la primera ola', 'MAX(SUM(SKUs!$D$6:$D$24),1)'),
 ('c_k2', 'Factor de escala de la segunda ola', 'p_n2*p_rel2/c_n1'),
 ('c_rs_unit', 'Costo total por registro (B/.; ITBMS sobre abogado y traducciones)', 'p_rs_off+p_rs_iea+(p_rs_law+p_rs_trad)*(1+p_itbms_srv)+p_rs_smp'),
 ('c_bid_m', 'Meses de colateral de la fianza de propuesta (entero)', 'ROUND(p_bid_m,0)'),
 ('c_n_reg1', 'Registros aprobados esperados, primera ola', 'SUM(SKUs!$D$6:$D$24)*p_rs_succ'),
 ('c_n_reg2', 'Registros aprobados esperados, segunda ola', 'p_n2*p_rs_succ'),
]
for nm, lab, f in calc:
    put(ws, r, 1, nm, NOTE); put(ws, r, 2, lab, BLACK); put(ws, r, 7, '=' + f, BLACK, '0.00' if nm == 'c_k2' else NUM)
    name(nm, f"Supuestos!$G${r}"); r += 1
ws.freeze_panes = 'C7'

# =============================================================== SKUs
ws = wb.create_sheet('SKUs')
cols = ['ID', 'Producto', 'Grupo', 'Incluir (1/0)', 'FOB base (B/./u)', 'Logística (B/./u)', 'Fuente del costo',
        'CSS LPU: unid./año', 'CSS LPU: precio', 'MINSA central: unid./año', 'MINSA central: precio',
        'CSS desc.: unid./año', 'CSS desc.: precio', 'MINSA/patronatos desc.: unid./año', 'MINSA/patronatos desc.: precio',
        'Costo puesto en Panamá', 'Precio oferta LPU', 'Precio oferta MINSA central', 'Precio oferta CSS desc.', 'Precio oferta MINSA desc.',
        'Margen LPU', 'Margen MINSA central', 'Margen CSS desc.', 'Margen MINSA desc.',
        'Oferta LPU', 'Oferta MINSA central', 'Oferta CSS desc.', 'Oferta MINSA desc.',
        'Ventas CSS desc. (madurez)', 'Ventas MINSA desc. (madurez)', 'Ventas LPU si gana', 'Ventas MINSA central si gana',
        'Costo CSS desc.', 'Costo MINSA desc.', 'Costo LPU', 'Costo MINSA central',
        'Ventas esperadas año maduro', 'Utilidad bruta esperada', 'Costo FOB (para Zywie)', 'Oferta LPU con fianza de propuesta (B/./año)']
put(ws, 1, 1, 'Portafolio y economía por producto', TITLE)
put(ws, 2, 1, 'Demanda pública anualizada (LPU 01-2026; subasta del MINSA de 2026 o compra conjunta de 2023 x p_cc_w; compras menores 2025-2026 con precio ponderado). Ventas a madurez antes de rampa y erosión.', NOTE)
put(ws, 3, 1, 'Canales: LPU = licitación de precio único de la CSS; MINSA central = subasta en reversa o compra conjunta; desc. = contratación menor y compra menor ágil. Costo = referencia FOB x índice x (1 + empaque) + logística.', NOTE)
hdr(ws, 5, range(1, len(cols) + 1), cols)
ws.row_dimensions[5].height = 60
for i, s in enumerate(SKUS):
    r = 6 + i
    q2v = f"={s['q2']}*p_cc_w" if (s.get('c2', 'cc') == 'cc' and s['q2'] > 0) else s['q2']
    vals = [s['id'], s['name'], s['group'], s['inc'], s['fob'], s['adder'], s['src_cost'], s['q1'], s['p1'], q2v, s['p2'], s['q3'], s['p3'], s['q4'], s['p4']]
    for j, v in enumerate(vals, start=1):
        fmt = NUM if j in (8, 10, 12, 14) else (NUM4 if j in (5, 6, 9, 11, 13, 15) else None)
        put(ws, r, j, v, BLUE if j >= 4 and j != 7 else BLACK, fmt, wrap=(j in (2, 7)))
    put(ws, r, 16, f'=E{r}*p_sup_idx*(1+p_pack)+F{r}*p_log_mult', BLACK, NUM4)
    put(ws, r, 17, f'=I{r}*p_pf_cen', BLACK, NUM4)
    put(ws, r, 18, f'=K{r}*p_pf_cen', BLACK, NUM4)
    put(ws, r, 19, f'=M{r}*p_pf_dec', BLACK, NUM4)
    put(ws, r, 20, f'=O{r}*p_pf_dec', BLACK, NUM4)
    for j, pc in enumerate(('Q', 'R', 'S', 'T')):
        put(ws, r, 21 + j, f'=IF({pc}{r}>0,1-P{r}/{pc}{r},-9)', BLACK, PCT)
    for j, (qc, gc, mn) in enumerate((('H', 'U', 'p_gm_cen'), ('J', 'V', 'p_gm_cen'), ('L', 'W', 'p_gm_dec'), ('N', 'X', 'p_gm_dec'))):
        put(ws, r, 25 + j, f'=IF(AND($D{r}=1,{qc}{r}>0,{gc}{r}>={mn}),1,0)', BLACK, '0')
    put(ws, r, 29, f'=AA{r}*L{r}*p_mkt*S{r}*p_rs_succ*p_share_dec', BLACK, NUM)
    put(ws, r, 30, f'=AB{r}*N{r}*p_mkt*T{r}*p_rs_succ*p_share_dec', BLACK, NUM)
    put(ws, r, 31, f'=Y{r}*H{r}*p_mkt*Q{r}*p_rs_succ*p_realiz', BLACK, NUM)
    put(ws, r, 32, f'=Z{r}*J{r}*p_mkt*R{r}*p_rs_succ*p_realiz_minsa', BLACK, NUM)
    put(ws, r, 33, f'=IF(S{r}>0,AC{r}*P{r}/S{r},0)', BLACK, NUM)
    put(ws, r, 34, f'=IF(T{r}>0,AD{r}*P{r}/T{r},0)', BLACK, NUM)
    put(ws, r, 35, f'=IF(Q{r}>0,AE{r}*P{r}/Q{r},0)', BLACK, NUM)
    put(ws, r, 36, f'=IF(R{r}>0,AF{r}*P{r}/R{r},0)', BLACK, NUM)
    put(ws, r, 37, f'=AC{r}+AD{r}+AE{r}*p_pwin_cen+AF{r}*p_pwin_cen*p_minsa_occ', BLACK, NUM)
    put(ws, r, 38, f'=(AC{r}-AG{r})+(AD{r}-AH{r})+(AE{r}-AI{r})*p_pwin_cen+(AF{r}-AJ{r})*p_pwin_cen*p_minsa_occ', BLACK, NUM)
    put(ws, r, 39, f'=IF(P{r}>0,E{r}*p_sup_idx*(1+p_pack)/P{r},0)', BLACK, PCT)
    put(ws, r, 40, f'=IF(AND(Y{r}=1,H{r}>0),H{r}*p_mkt*Q{r}*p_rs_succ,0)', BLACK, NUM)
last = 6 + len(SKUS) - 1
tr = last + 1
put(ws, tr, 2, 'Totales (productos incluidos)', BOLD)
for c in range(29, 39):
    L = CL(c)
    put(ws, tr, c, f'=SUM({L}6:{L}{last})', BOLD, NUM)
put(ws, tr, 39, f'=IF(SUM(AG6:AJ{last})>0,SUMPRODUCT(AG6:AG{last}+AH6:AH{last}+AI6:AI{last}+AJ6:AJ{last},AM6:AM{last})/SUM(AG6:AJ{last}),0)', BOLD, PCT)
name('s_Rd', f"SKUs!$AC${tr}:$AD${tr}")
put(ws, tr, 40, f'=SUM(AN6:AN{last})', BOLD, NUM)
for nm, c in (('s_Rd3', 'AC'), ('s_Rd4', 'AD'), ('s_Rc1', 'AE'), ('s_Rc2', 'AF'), ('s_Rd3k', 'AG'), ('s_Rd4k', 'AH'), ('s_Rc1k', 'AI'), ('s_Rc2k', 'AJ'), ('s_fob_share', 'AM'), ('s_rev_exp', 'AK'), ('s_gp_exp', 'AL'), ('s_bid', 'AN')):
    name(nm, f"SKUs!${c}${tr}")
setw(ws, {'A': 8, 'B': 40, 'C': 8, 'D': 8, 'G': 34})
for c in range(5, 41):
    if CL(c) != 'G':
        ws.column_dimensions[CL(c)].width = 11
ws.freeze_panes = 'C6'

# =============================================================== Mensual
ws = wb.create_sheet('Mensual')
setw(ws, {'A': 46, 'B': 9})
put(ws, 1, 1, 'Flujo mensual (B/.)', TITLE)
M0 = 3  # column C = month 0; month m -> column 3+m
def mc_(m): return CL(3 + m)
R = {}
rows = ['mes', 'anio', 'eros', 'activo', 'ramp1', 'ramp2', 'cen_on', 'cen2_on', 'cen2m_on', 'mult_dec', 'mult_cen', 'mult_cen2', 'v_dec', 'v_cen', 'ventas',
        'k_dec', 'k_cen', 'costo', 'inv', 'llegadas', 'pago_ant', 'pago_sal', 'pagos', 'cobranza', 'var', 'fijos', 'posreg', 'setup',
        'registros', 'anuales', 'afianzado', 'fianza_prop', 'colateral', 'dcol', 'prima', 'factoring', 'perdida_cierre', 'uai', 'imp_pagado',
        'aviso', 'caja_pre', 'caja_min', 'aporte', 'dividendo', 'devolucion', 'dev_imp', 'caja', 'flujo_inv']
labels = {'mes': 'Mes', 'anio': 'Año', 'eros': 'Factor de erosión de precios', 'activo': 'Negocio activo (1/0)',
          'ramp1': 'Rampa primera ola', 'ramp2': 'Rampa segunda ola', 'cen_on': 'Contratos centralizados activos', 'cen2_on': 'LPU con segunda ola',
          'cen2m_on': 'MINSA central con segunda ola', 'mult_dec': 'Multiplicador descentralizado (olas)', 'mult_cen': 'Multiplicador LPU (olas)',
          'mult_cen2': 'Multiplicador MINSA central (olas)',
          'v_dec': 'Ventas descentralizadas', 'v_cen': 'Ventas centralizadas', 'ventas': 'VENTAS TOTALES',
          'k_dec': 'Costo de ventas descentralizado', 'k_cen': 'Costo de ventas centralizado', 'costo': 'COSTO DE VENTAS',
          'inv': 'Inventario objetivo al cierre (costo)', 'llegadas': 'Mercancía recibida (costo, incl. mermas)',
          'pago_ant': 'Pago de anticipos al proveedor', 'pago_sal': 'Pago de saldos al proveedor', 'pagos': 'PAGOS AL PROVEEDOR',
          'cobranza': 'COBRANZA (incl. liquidación de inventario al cierre)', 'var': 'Gastos variables (distribución y multas)',
          'fijos': 'Costos fijos', 'posreg': 'Análisis posregistro', 'setup': 'Gastos de arranque', 'registros': 'Registros sanitarios',
          'anuales': 'Cargos anuales', 'afianzado': 'Monto afianzado (cumplimiento, sobre la cantidad máxima)',
          'fianza_prop': 'Fianza de propuesta vigente (LPU)', 'aviso': 'Impuesto de aviso de operación pagado', 'colateral': 'Colateral inmovilizado',
          'dcol': 'Cambio en colateral', 'prima': 'Prima de fianzas', 'factoring': 'Costo de factoring',
          'perdida_cierre': 'Pérdida de inventario por cierre', 'uai': 'UTILIDAD ANTES DE IMPUESTOS', 'imp_pagado': 'Impuesto sobre la renta pagado',
          'caja_pre': 'Caja antes de aportes', 'caja_min': 'Caja mínima requerida', 'aporte': 'APORTE DE CAPITAL', 'dividendo': 'DIVIDENDO PAGADO',
          'devolucion': 'Devolución de caja a los socios después de un cierre', 'dev_imp': 'Impuesto de dividendos sobre lo devuelto por encima de lo aportado',
          'caja': 'CAJA AL CIERRE', 'flujo_inv': 'FLUJO DE LOS ACCIONISTAS (sin valor terminal)'}
for i, k in enumerate(rows):
    R[k] = 3 + i
    put(ws, R[k], 1, labels[k], BOLD if labels[k].isupper() or k in ('ventas', 'costo', 'uai', 'aporte', 'caja') else BLACK)
put(ws, 2, 3, 'Mes 0', NOTE)
rng60 = lambda k: f"$D${R[k]}:$BK${R[k]}"
name('m_mes', f"Mensual!$D${R['mes']}:$BK${R['mes']}")
for m in range(0, 62):
    c = 3 + m; L = CL(c); Lp = CL(c - 1); Ln = CL(c + 1)
    f = {}
    f['mes'] = m
    if m == 0:
        put(ws, R['mes'], c, 0, BLACK, '0')
        put(ws, R['flujo_inv'], c, '=-c_phase0', BLACK, NUM)
        put(ws, R['caja'], c, 0, BLACK, NUM)
        put(ws, R['inv'], c, 0, BLACK, NUM)
        put(ws, R['colateral'], c, 0, BLACK, NUM)
        continue
    put(ws, R['mes'], c, m, BOLD, '0')
    mm = f'{L}${R["mes"]}'
    put(ws, R['anio'], c, f'=INT((MIN({mm},60)-1)/12)+1', BLACK, '0')
    put(ws, R['eros'], c, f'=(1-p_erosion)^({L}{R["anio"]}-1)', BLACK, '0.000')
    put(ws, R['activo'], c, f'=IF({mm}<=ROUND(p_stop_det,0),1,0)', BLACK, '0')
    put(ws, R['ramp1'], c, f'=MAX(0,MIN(1,({mm}-c_first_dec+1)/p_ramp_m))', BLACK, '0.00')
    put(ws, R['ramp2'], c, f'=MAX(0,MIN(1,({mm}-c_first_dec2+1)/p_ramp_m))', BLACK, '0.00')
    put(ws, R['cen_on'], c, f'=IF({mm}>=c_cen_start,1,0)', BLACK, '0')
    put(ws, R['cen2_on'], c, f'=IF({mm}>=c_c2_lpu,1,0)', BLACK, '0')
    put(ws, R['cen2m_on'], c, f'=IF({mm}>=c_c2_minsa,1,0)', BLACK, '0')
    put(ws, R['mult_dec'], c, f'=1+c_k2*{L}{R["ramp2"]}', BLACK, '0.00')
    put(ws, R['mult_cen'], c, f'=1+c_k2*{L}{R["cen2_on"]}', BLACK, '0.00')
    put(ws, R['mult_cen2'], c, f'=1+c_k2*{L}{R["cen2m_on"]}', BLACK, '0.00')
    put(ws, R['v_dec'], c, f'=(s_Rd3+s_Rd4)/12*{L}{R["ramp1"]}*{L}{R["mult_dec"]}*{L}{R["eros"]}*{L}{R["activo"]}', BLACK, NUM)
    put(ws, R['v_cen'], c, f'=(s_Rc1*p_pwin_cen*{L}{R["mult_cen"]}+s_Rc2*p_pwin_cen*p_minsa_occ*{L}{R["mult_cen2"]})/12*{L}{R["cen_on"]}*{L}{R["eros"]}*{L}{R["activo"]}', BLACK, NUM)
    put(ws, R['ventas'], c, f'={L}{R["v_dec"]}+{L}{R["v_cen"]}', BOLD, NUM)
    put(ws, R['k_dec'], c, f'=(s_Rd3k+s_Rd4k)/12*{L}{R["ramp1"]}*{L}{R["mult_dec"]}*{L}{R["activo"]}', BLACK, NUM)
    put(ws, R['k_cen'], c, f'=(s_Rc1k*p_pwin_cen*{L}{R["mult_cen"]}+s_Rc2k*p_pwin_cen*p_minsa_occ*{L}{R["mult_cen2"]})/12*{L}{R["cen_on"]}*{L}{R["activo"]}', BLACK, NUM)
    put(ws, R['costo'], c, f'={L}{R["k_dec"]}+{L}{R["k_cen"]}', BOLD, NUM)
    if m == 61:
        for k in ('inv', 'llegadas'):
            pass
        continue
    put(ws, R['inv'], c, (f'=IF(AND({mm}>=c_first_dec,{mm}<=ROUND(p_stop_det,0)),(s_Rd3k+s_Rd4k)/12*{Ln}{R["ramp1"]}*{Ln}{R["mult_dec"]}*p_cover_m'
                          f'+(s_Rc1k*p_pwin_cen*{Ln}{R["mult_cen"]}+s_Rc2k*p_pwin_cen*p_minsa_occ*{Ln}{R["mult_cen2"]})/12*{Ln}{R["cen_on"]}*p_cen_stock_m,0)'), BLACK, NUM)
    put(ws, R['llegadas'], c, f'=IF({mm}>ROUND(p_stop_det,0),0,{L}{R["costo"]}*(1+p_expiry)+{L}{R["inv"]}-{Lp}{R["inv"]})', BLACK, NUM)
    put(ws, R['pago_ant'], c, f'=p_adv*IFERROR(INDEX({rng60("llegadas")},1,{mm}+c_lead_m),0)', BLACK, NUM)
    put(ws, R['pago_sal'], c, f'=IF(AND({mm}+c_ship_m-c_credit_m>=1,{mm}+c_ship_m-c_credit_m<=60),(1-p_adv)*INDEX({rng60("llegadas")},1,{mm}+c_ship_m-c_credit_m),0)', BLACK, NUM)
    put(ws, R['pagos'], c, f'={L}{R["pago_ant"]}+{L}{R["pago_sal"]}', BOLD, NUM)
    put(ws, R['cobranza'], c, (f'=(1-p_fact_share)*IF({mm}-c_dso_m>=1,INDEX({rng60("ventas")},1,{mm}-c_dso_m),0)*(1-p_baddebt)'
                               f'+p_fact_share*IF({mm}-1>=1,INDEX({rng60("ventas")},1,{mm}-1),0)*(1-p_baddebt)*(1-p_fact_cost*c_dso_m)'
                               f'+IF({mm}=ROUND(p_stop_det,0)+1,INDEX({rng60("inv")},1,ROUND(p_stop_det,0))*p_inv_liq,0)'), BOLD, NUM)
    put(ws, R['var'], c, f'=(p_dist+p_penal)*{L}{R["ventas"]}', BLACK, NUM)
    put(ws, R['fijos'], c, f'=IF({mm}>ROUND(p_stop_det,0)+3,0,SUMPRODUCT(f_used*({mm}>=f_start)))', BLACK, NUM)
    put(ws, R['posreg'], c, f'=IF({mm}>ROUND(p_stop_det,0)+3,0,c_n_reg1*p_postreg/12*IF({mm}>=c_reg_m,1,0)+c_n_reg2*p_postreg/12*IF({mm}>=c_m2+c_rt,1,0))', BLACK, NUM)
    put(ws, R['setup'], c, f'=SUMIF(u_month,{mm},u_used)', BLACK, NUM)
    put(ws, R['registros'], c, f'=IF({mm}=c_rs_pay_m,SUM(SKUs!$D$6:$D$24)*c_rs_unit,0)+IF({mm}=MAX(1,MIN(60,c_m2)),p_n2*c_rs_unit,0)', BLACK, NUM)
    put(ws, R['anuales'], c, f'=IF({mm}>ROUND(p_stop_det,0)+3,0,IF(OR({mm}=1,{mm}=13,{mm}=25,{mm}=37,{mm}=49),c_annual,0))', BLACK, NUM)
    ov1 = f'IF(AND({mm}-c_cen_start>=15,MOD({mm}-c_cen_start,15)<ROUND(p_bond_tail_css,0)),1,0)'
    ov2 = f'IF(AND({mm}-c_cen_start>=12,MOD({mm}-c_cen_start,12)<ROUND(p_bond_tail_minsa,0)),1,0)'
    put(ws, R['afianzado'], c, (f'=(s_Rc1/p_realiz*p_pwin_cen*p_bond_css*15/12*{L}{R["mult_cen"]}*(1+{ov1})'
                                f'+s_Rc2/p_realiz_minsa*p_pwin_cen*p_minsa_occ*p_bond_minsa*{L}{R["mult_cen2"]}*(1+{ov2}))*{L}{R["cen_on"]}*{L}{R["activo"]}'), BLACK, NUM)
    put(ws, R['fianza_prop'], c, f'=p_bid_bond*s_bid*15/12*{L}{R["mult_cen"]}*IF(AND({mm}>=c_cen_start-c_bid_m,MOD({mm}-(c_cen_start-c_bid_m),15)<c_bid_m),1,0)*{L}{R["activo"]}', BLACK, NUM)
    put(ws, R['colateral'], c, f'=({L}{R["afianzado"]}+{L}{R["fianza_prop"]})*p_collat', BLACK, NUM)
    put(ws, R['dcol'], c, f'={L}{R["colateral"]}-{Lp}{R["colateral"]}', BLACK, NUM)
    put(ws, R['prima'], c, f'=({L}{R["afianzado"]}+{L}{R["fianza_prop"]})*p_bond_prem/12', BLACK, NUM)
    put(ws, R['factoring'], c, f'=p_fact_share*{L}{R["ventas"]}*(1-p_baddebt)*p_fact_cost*c_dso_m', BLACK, NUM)
    put(ws, R['perdida_cierre'], c, f'=IF({mm}=ROUND(p_stop_det,0)+1,INDEX({rng60("inv")},1,ROUND(p_stop_det,0))*(1-p_inv_liq),0)', BLACK, NUM)
    put(ws, R['uai'], c, (f'={L}{R["ventas"]}-{L}{R["costo"]}*(1+p_expiry)-{L}{R["var"]}-{L}{R["ventas"]}*p_baddebt-{L}{R["factoring"]}'
                          f'-{L}{R["fijos"]}-{L}{R["posreg"]}-{L}{R["setup"]}-{L}{R["registros"]}-{L}{R["anuales"]}-{L}{R["prima"]}-{L}{R["perdida_cierre"]}'), BOLD, NUM)
    tparts = []
    if m in (15, 27, 39, 51):
        yy_ = (m - 3) // 12                       # 1..4: balance of year yy_ = its tax minus the installments paid on it
        tcol = CL(2 + yy_); pcol = CL(1 + yy_)
        tparts.append(f'Anual!${tcol}$14' + (f'-Anual!${pcol}$14' if yy_ >= 2 else ''))
    if m >= 13 and (m - 12 * ((m - 1) // 12)) in (6, 9, 12):
        yy_ = (m - 1) // 12                       # installment during year yy_+1 on the tax of year yy_
        tparts.append(f'Anual!${CL(2 + yy_)}$14/3')
    put(ws, R['imp_pagado'], c, '=' + ('+'.join(tparts) if tparts else '0'), GREEN, NUM)
    if m in (15, 27, 39, 51):
        yy = (m - 3) // 12; c12 = CL(3 + 12 * yy); ycl = CL(2 + yy)
        put(ws, R['aviso'], c, (f'=IF({mm}<=ROUND(p_stop_det,0)+3,MIN(60000,MAX(100,p_aviso*MAX(0,SUM($D${R["aporte"]}:${c12}${R["aporte"]})'
                                f'+SUM(Anual!$C$15:${ycl}$15)-SUM($D${R["dividendo"]}:${c12}${R["dividendo"]})-SUM($D${R["devolucion"]}:${c12}${R["devolucion"]})))),0)'), BLACK, NUM)
    else:
        put(ws, R['aviso'], c, 0, BLACK, NUM)
    put(ws, R['caja_pre'], c, (f'={Lp}{R["caja"]}+{L}{R["cobranza"]}-{L}{R["pagos"]}-{L}{R["var"]}-{L}{R["fijos"]}-{L}{R["posreg"]}-{L}{R["setup"]}'
                               f'-{L}{R["registros"]}-{L}{R["anuales"]}-{L}{R["dcol"]}-{L}{R["prima"]}-{L}{R["imp_pagado"]}-{L}{R["aviso"]}'), BLACK, NUM)
    put(ws, R['caja_min'], c, f'=p_mincash_m*{L}{R["fijos"]}', BLACK, NUM)
    put(ws, R['aporte'], c, f'=MAX(0,{L}{R["caja_min"]}-{L}{R["caja_pre"]})', BOLD, NUM)
    if m in (24, 36, 48):
        y = m // 12
        ycol = CL(2 + y)
        put(ws, R['dividendo'], c, (f'=p_payout*MAX(0,MIN({L}{R["caja_pre"]}+{L}{R["aporte"]}-2*{L}{R["caja_min"]}-Anual!${ycol}$14,'
                                     f'SUM(Anual!$C$15:${ycol}$15)-SUM($D${R["dividendo"]}:{Lp}${R["dividendo"]})))'), BOLD, NUM)
    else:
        put(ws, R['dividendo'], c, 0, BLACK, NUM)
    put(ws, R['devolucion'], c, f'=IF(AND({mm}>ROUND(p_stop_det,0)+3,{mm}<60),MAX(0,{L}{R["caja_pre"]}+{L}{R["aporte"]}-{L}{R["dividendo"]}),0)', BLACK, NUM)
    put(ws, R['dev_imp'], c, (f'=p_divtax*(MAX(0,SUM($D${R["devolucion"]}:{L}{R["devolucion"]})-SUM($D${R["aporte"]}:{L}{R["aporte"]}))'
                              f'-MAX(0,SUM($C${R["devolucion"]}:{Lp}{R["devolucion"]})-SUM($D${R["aporte"]}:{L}{R["aporte"]})))'), BLACK, NUM)
    put(ws, R['caja'], c, f'={L}{R["caja_pre"]}+{L}{R["aporte"]}-{L}{R["dividendo"]}-{L}{R["devolucion"]}', BOLD, NUM)
    put(ws, R['flujo_inv'], c, f'=-{L}{R["aporte"]}+{L}{R["dividendo"]}*(1-p_divtax)+{L}{R["devolucion"]}-{L}{R["dev_imp"]}', BLACK, NUM)
for c in range(3, 65):
    ws.column_dimensions[CL(c)].width = 10
ws.freeze_panes = 'C4'
for k in ('mes', 'ventas', 'costo', 'uai', 'aporte', 'caja'):
    for c in range(1, 65):
        ws.cell(row=R[k], column=c).fill = GREY
put(ws, R['flujo_inv'] + 2, 1, 'La columna del mes 61 solo sirve para mirar el costo del mes siguiente al fijar el inventario del mes 60.', NOTE)

# =============================================================== Anual
ws = wb.create_sheet('Anual')
setw(ws, {'A': 4, 'B': 52, 'C': 13, 'D': 13, 'E': 13, 'F': 13, 'G': 13, 'H': 13})
put(ws, 1, 2, 'Estado de resultados e impuestos (B/.)', TITLE)
hdr(ws, 3, range(2, 9), ['Concepto', 'Año 1', 'Año 2', 'Año 3', 'Año 4', 'Año 5', 'Total'])
lines = [
 (4, 'Ventas', 'ventas', 1),
 (5, 'Costo de ventas (incl. mermas)', None, -1),
 (6, 'Gastos variables, incobrables y factoring', None, -1),
 (7, 'Costos fijos y cargos anuales', None, -1),
 (8, 'Arranque y registros sanitarios', None, -1),
 (9, 'Fianzas, posregistro y pérdidas por cierre', None, -1),
 (10, 'Utilidad antes de impuestos', 'uai', 1),
]
def ysum(row_key, y):
    a = CL(3 + 12 * (y - 1) + 1); b = CL(3 + 12 * y)
    return f"SUM(Mensual!{a}{R[row_key]}:{b}{R[row_key]})"
for y in range(1, 6):
    c = 2 + y
    a = CL(3 + 12 * (y - 1) + 1); b = CL(3 + 12 * y)
    put(ws, 4, c, '=' + ysum('ventas', y), GREEN, NUM)
    put(ws, 5, c, f'=-{ysum("costo", y)}*(1+p_expiry)', GREEN, NUM)
    put(ws, 6, c, f'=-({ysum("var", y)}+{ysum("ventas", y)}*p_baddebt+{ysum("factoring", y)})', GREEN, NUM)
    put(ws, 7, c, f'=-({ysum("fijos", y)}+{ysum("anuales", y)})', GREEN, NUM)
    put(ws, 8, c, f'=-({ysum("setup", y)}+{ysum("registros", y)})', GREEN, NUM)
    put(ws, 9, c, f'=-({ysum("prima", y)}+{ysum("posreg", y)}+{ysum("perdida_cierre", y)})', GREEN, NUM)
    put(ws, 10, c, '=' + ysum('uai', y), GREEN, NUM)
    # loss pool: 20% of each prior loss within 5 years
    prior = '+'.join([f'0.2*MAX(0,-{CL(2 + k)}10)' for k in range(1, y)]) or '0'
    put(ws, 11, c, f'={prior}', BLACK, NUM)
    put(ws, 12, c, f'=MIN(0.5*MAX(0,{CL(c)}10),{CL(c)}11)', BLACK, NUM)
    put(ws, 13, c, f'=MAX(0,{CL(c)}10)-{CL(c)}12', BLACK, NUM)
    put(ws, 14, c, f'=MAX(p_isr*{CL(c)}13,IF({CL(c)}4>1500000,0.011675*{CL(c)}4,0))', BLACK, NUM)
    put(ws, 15, c, f'={CL(c)}10-{CL(c)}14', BOLD, NUM)
    put(ws, 16, c, f'=IF({CL(c)}4>0,{CL(c)}15/{CL(c)}4,0)', BLACK, PCT)
for y in range(1, 6):
    put(ws, 17, 2 + y, '=' + ysum('aviso', y), GREEN, NUM)
put(ws, 17, 2, 'Memo: impuesto de aviso de operación pagado (no se deduce de la renta en el modelo)', NOTE)
put(ws, 17, 8, '=SUM(C17:G17)', BOLD, NUM)
labs = {11: 'Pérdidas acumuladas deducibles este año (20% de cada pérdida previa)', 12: 'Deducción aplicada (tope 50% de la renta)',
        13: 'Renta gravable', 14: 'Impuesto sobre la renta (o CAIR si ventas > B/.1.5 millones)', 15: 'Utilidad neta', 16: 'Margen neto'}
for r_, lab in [(l[0], l[1]) for l in lines] + list(labs.items()):
    put(ws, r_, 2, lab, BOLD if r_ in (4, 10, 15) else BLACK)
for r_ in range(4, 16):
    if r_ in (11, 12):
        continue
    put(ws, r_, 8, f'=SUM(C{r_}:G{r_})', BOLD, NUM)
put(ws, 18, 2, 'Punto de equilibrio (ventas anuales que cubren fijos)', SUB)
put(ws, 19, 2, 'Costos fijos anuales a plena operación', BLACK); put(ws, 19, 3, '=f_total*12+c_annual+c_n_reg1*p_postreg', BLACK, NUM)
put(ws, 20, 2, 'Margen bruto promedio del portafolio', BLACK); put(ws, 20, 3, '=IF(s_rev_exp>0,s_gp_exp/s_rev_exp,0)', BLACK, PCT)
put(ws, 21, 2, 'Gastos variables (% ventas)', BLACK); put(ws, 21, 3, '=p_dist+p_penal+p_baddebt+p_expiry*(1-C20)', BLACK, PCT)
put(ws, 22, 2, 'Ventas de equilibrio (B/./año)', BOLD); put(ws, 22, 3, '=IF(C20-C21>0,C19/(C20-C21),0)', BOLD, NUM)
put(ws, 23, 2, 'Ventas esperadas de la primera ola en madurez (B/./año)', BLACK); put(ws, 23, 3, '=s_rev_exp', GREEN, NUM)

# =============================================================== Valoración
ws = wb.create_sheet('Valoración')
setw(ws, {'A': 4, 'B': 62, 'C': 16, 'D': 60})
put(ws, 1, 2, 'Valoración y capital requerido', TITLE)
put(ws, 2, 2, '=CONCATENATE("Escenario activo: ",Supuestos!C3,IF(p_sl=1," | tarifas SignatureLex"," | tarifas de mercado"))', BOLD)
last_col = CL(63)
rows_v = [
 (4, 'Capital aportado por los socios durante 60 meses', f'=SUM(Mensual!D{R["aporte"]}:{last_col}{R["aporte"]})', NUM),
 (5, 'Gasto de fase 0 (antes de constituir, lo pagan los socios)', '=c_phase0', NUM),
 (6, 'Dividendos pagados (antes del 10%)', f'=SUM(Mensual!D{R["dividendo"]}:{last_col}{R["dividendo"]})', NUM),
 (7, 'Caja al mes 60', f'=Mensual!{last_col}{R["caja"]}', NUM),
 (8, 'Cuentas por cobrar al mes 60', f'=SUMPRODUCT((m_mes>60-c_dso_m)*Mensual!$D${R["ventas"]}:$BK${R["ventas"]})*(1-p_baddebt)*(1-p_fact_share)+p_fact_share*Mensual!{last_col}{R["ventas"]}*(1-p_baddebt)*(1-p_fact_cost*c_dso_m)', NUM),
 (9, 'Inventario al mes 60', f'=Mensual!{last_col}{R["inv"]}', NUM),
 (10, 'Colateral de fianzas al mes 60', f'=Mensual!{last_col}{R["colateral"]}', NUM),
 (11, 'Cuentas por pagar al proveedor al mes 60', f'=SUMPRODUCT((m_mes>60-MAX(0,c_credit_m-c_ship_m))*Mensual!$D${R["llegadas"]}:$BK${R["llegadas"]})*(1-p_adv)', NUM),
 (12, 'Capital de trabajo neto al mes 60', '=C8+C9+C10-C11', NUM),
 (13, 'Impuesto sobre la renta (saldo después de las cuotas estimadas) y aviso de operación del año 5 por pagar', f'=Anual!G14-Anual!F14+IF(60<=ROUND(p_stop_det,0)+3,MIN(60000,MAX(100,p_aviso*MAX(0,C4+SUM(Anual!C15:G15)-C6-SUM(Mensual!D{R["devolucion"]}:{last_col}{R["devolucion"]})))),0)', NUM),
 (14, 'Valor terminal conservador (caja + recuperación del capital de trabajo)', '=C7+p_recov*C12-C13', NUM),
 (15, 'Valor terminal como negocio en marcha (caja + múltiplo de la utilidad del año 5)', '=C7-C13+p_gc_mult*MAX(0,Anual!G15)', NUM),
 (16, 'Valor terminal usado en el caso "en marcha" (el mayor)', '=MAX(C14,C15)', NUM),
 (17, 'Valor terminal neto del impuesto de dividendos (conservador)', f'=C14-p_divtax*MAX(0,C14-MAX(0,C4-SUM(Mensual!D{R["devolucion"]}:{last_col}{R["devolucion"]})))', NUM),
 (18, 'Valor terminal neto del impuesto de dividendos (en marcha)', f'=C16-p_divtax*MAX(0,C16-MAX(0,C4-SUM(Mensual!D{R["devolucion"]}:{last_col}{R["devolucion"]})))', NUM),
]
for r_, lab, f, fmt in rows_v:
    put(ws, r_, 2, lab, BLACK); put(ws, r_, 3, f, GREEN if 'Mensual' in f or 'Anual' in f else BLACK, fmt)
# investor flow rows (month 0..60)
put(ws, 21, 2, 'Flujo de los accionistas, conservador (mes 0 a 60, sigue a la derecha)', SUB)
put(ws, 22, 2, 'Flujo de los accionistas, en marcha', SUB)
for m in range(0, 61):
    c = 3 + m
    L = CL(c)
    src = f'Mensual!{CL(3 + m)}{R["flujo_inv"]}'
    put(ws, 20, c, m, BOLD, '0')
    put(ws, 21, c, f'={src}' + ('+$C$17' if m == 60 else ''), GREEN, NUM)
    put(ws, 22, c, f'={src}' + ('+$C$18' if m == 60 else ''), GREEN, NUM)
    ws.column_dimensions[L].width = 11 if c > 3 else 16
fl = f'$C$21:${CL(63)}$21'; fg = f'$C$22:${CL(63)}$22'
res_rows = [
 (24, 'VPN al 15% anual (conservador)', f'=SUMPRODUCT({fl},1/(1+0.15)^($C$20:${CL(63)}$20/12))', NUM),
 (25, 'VPN a la tasa de descuento del escenario (conservador)', f'=SUMPRODUCT({fl},1/(1+p_disc)^($C$20:${CL(63)}$20/12))', NUM),
 (26, 'VPN al 15% anual (en marcha)', f'=SUMPRODUCT({fg},1/(1+0.15)^($C$20:${CL(63)}$20/12))', NUM),
 (27, 'TIR anual (conservador)', f'=IFERROR((1+IRR({fl},0.01))^12-1,IFERROR((1+IRR({fl},-0.1))^12-1,IFERROR((1+IRR({fl},-0.19))^12-1,IFERROR((1+IRR({fl},-0.3))^12-1,"n.d."))))', PCT),
 (28, 'TIR anual (en marcha)', f'=IFERROR((1+IRR({fg},0.01))^12-1,IFERROR((1+IRR({fg},-0.1))^12-1,IFERROR((1+IRR({fg},-0.19))^12-1,"n.d.")))', PCT),
 (29, 'Múltiplo del capital (lo que vuelve / lo que se pone), conservador', f'=IF(-SUMIF({fl},"<0")>0,SUMIF({fl},">0")/-SUMIF({fl},"<0"),0)', MULT),
 (30, 'Capital máximo acumulado (pico de necesidad de caja de los socios)', f'=-MIN(C34:{CL(63)}34)', NUM),
 (31, 'Ganancia de Zywie por su margen sobre el precio de fábrica (60 meses)', f'=SUM(Mensual!D{R["llegadas"]}:{CL(63)}{R["llegadas"]})*s_fob_share*p_zy_markup/(1+p_zy_markup)', NUM),
]
for r_, lab, f, fmt in res_rows:
    put(ws, r_, 2, lab, BOLD); c = put(ws, r_, 3, f, BLACK, fmt); c.fill = YEL
put(ws, 33, 2, 'Flujo acumulado (conservador, sin valor terminal)', BLACK)
put(ws, 34, 2, 'Flujo acumulado de los accionistas', BLACK)
for m in range(0, 61):
    c = 3 + m; L = CL(c)
    put(ws, 34, c, f'=SUM($C$21:{L}21)-IF({m}=60,$C$17,0)', BLACK, NUM)

put(ws, 36, 2, 'Vista de un socio panameño (participación y aportes según el trato)', SUB)
put(ws, 37, 2, 'Participación del socio en las utilidades', BLACK); put(ws, 37, 3, '=(1-p_patel_eq)*p_jaime_frac', BLACK, PCT)
put(ws, 38, 2, 'Porcentaje del capital que aporta el socio', BLACK); put(ws, 38, 3, '=(1-p_patel_cash)*p_jaime_frac', BLACK, PCT)
put(ws, 39, 2, 'Capital total que pondría el socio (incluye su tercio de la fase 0)', BOLD); put(ws, 39, 3, '=C4*C38+C5*p_jaime_frac', BOLD, NUM)
put(ws, 40, 2, 'Lo que recibiría el socio (dividendos y valor terminal netos, conservador)', BOLD)
put(ws, 40, 3, f'=(C6*(1-p_divtax)+SUM(Mensual!D{R["devolucion"]}:{last_col}{R["devolucion"]})-SUM(Mensual!D{R["dev_imp"]}:{last_col}{R["dev_imp"]})+C17)*C37', BOLD, NUM)
put(ws, 41, 2, 'Resultado neto del socio en 5 años (sin descontar)', BOLD); c = put(ws, 41, 3, '=C40-C39', BOLD, NUM); c.fill = YEL
put(ws, 44, 2, 'Reparto del valor económico entre Rutul Patel y el lado panameño (nominal, 60 meses, conservador)', SUB)
put(ws, 45, 2, 'Valor neto para todos los accionistas (lo que vuelve menos lo que se pone)', BLACK); put(ws, 45, 3, f'=SUM({fl})', BLACK, NUM)
put(ws, 46, 2, 'Rutul Patel: margen de Zywie + su parte de lo distribuido - su parte del capital', BLACK)
put(ws, 46, 3, f'=C31+p_patel_eq*SUMIF(Valoración!$D$21:${CL(63)}$21,">0")-p_patel_cash*C4', BLACK, NUM)
put(ws, 47, 2, 'Lado panameño (3 socios): su parte de lo distribuido - su parte del capital - fase 0', BLACK)
put(ws, 47, 3, f'=(1-p_patel_eq)*SUMIF(Valoración!$D$21:${CL(63)}$21,">0")-(1-p_patel_cash)*C4-C5', BLACK, NUM)
put(ws, 48, 2, 'Porcentaje del valor total que se lleva Rutul Patel', BOLD); put(ws, 48, 3, '=IF(C46+C47<>0,C46/(C46+C47),0)', BOLD, PCT)
put(ws, 49, 2, 'Nota: si el valor neto total es negativo, el porcentaje no tiene interpretación útil; mire los montos.', NOTE)

# =============================================================== Escenarios (fijo)
ws = wb.create_sheet('Escenarios')
setw(ws, {'A': 4, 'B': 60, 'C': 14, 'D': 14, 'E': 14, 'F': 14, 'G': 14, 'H': 14})
put(ws, 1, 2, 'Resultados por escenario (valores fijos del motor en Python)', TITLE)
put(ws, 2, 2, 'Para recalcular en vivo, cambie Supuestos!B3. Estos valores sirven de referencia y de control cruzado.', NOTE)
hdr(ws, 4, range(2, 6), ['Indicador', 'Pesimista', 'Base', 'Optimista'])
D_ = A['det']
ind = [('Ventas año 1', lambda d: d['revY'][0], NUM), ('Ventas año 2', lambda d: d['revY'][1], NUM), ('Ventas año 3', lambda d: d['revY'][2], NUM),
       ('Ventas año 5', lambda d: d['revY'][4], NUM), ('Utilidad neta año 1', lambda d: d['NIY'][0], NUM), ('Utilidad neta año 3', lambda d: d['NIY'][2], NUM),
       ('Utilidad neta año 5', lambda d: d['NIY'][4], NUM), ('Capital (aportes más fase 0)', lambda d: d['capital'], NUM),
       ('Mes del pico de capital', lambda d: d['peak_m'], '0'), ('VPN al 15% (conservador)', lambda d: d['npv15'], NUM),
       ('VPN al 30% (conservador)', lambda d: d['npv30'], NUM), ('TIR (conservador)', lambda d: d['irr'], PCT),
       ('Múltiplo del capital (conservador)', lambda d: d['moic'], MULT), ('Mes de la primera venta', lambda d: d['first_sale_m'], '0'),
       ('Mes del primer contrato centralizado', lambda d: d['cen_m'], '0'), ('Ganancia de Zywie por su margen', lambda d: d['zywie'], NUM)]
for i, (lab, fn, fmt) in enumerate(ind):
    r_ = 5 + i
    put(ws, r_, 2, lab, BLACK)
    for j, sc in enumerate(('pes', 'base', 'opt')):
        v = fn(D_[sc]); put(ws, r_, 3 + j, v if v == v else 'n.d.', BLACK, fmt)
r_ = 5 + len(ind) + 2
put(ws, r_, 2, 'Qué pasaría si (cada fila cambia el caso base)', SUB); r_ += 1
hdr(ws, r_, range(2, 8), ['Cambio', 'Capital', 'VPN 15%', 'TIR', 'Múltiplo', 'Ventas año 3'])
for w in A['whatif']:
    r_ += 1
    put(ws, r_, 2, w['tag'], BLACK, wrap=True)
    for j, (k, fmt) in enumerate((('cap', NUM), ('npv15', NUM), ('irr', PCT), ('moic', MULT), ('rev3', NUM))):
        v = w[k]; put(ws, r_, 3 + j, v if v == v else 'n.d.', BLACK, fmt)

# =============================================================== Simulación (fija)
ws = wb.create_sheet('Simulación')
setw(ws, {'A': 4, 'B': 58, 'C': 13, 'D': 13, 'E': 13, 'F': 13, 'G': 13, 'H': 13, 'I': 13, 'J': 13, 'K': 13})
put(ws, 1, 2, 'Simulación de Monte Carlo: 20,000 escenarios por estructura (mismas semillas)', TITLE)
put(ws, 2, 2, 'Cada escenario sortea los supuestos (triangular pesimista-base-optimista), los registros aprobados, cada renglón centralizado, fallas de planta o de suministro y las compuertas de decisión.', NOTE)
hdr(ws, 4, range(2, 12), ['Estructura', 'Se detiene en la fase 0', 'Cierra en la compuerta operativa', 'Falla de planta o suministro',
                          'Prob. de perder dinero', 'Prob. de perder más que la fase 0', 'VPN 15% promedio', 'Múltiplo mediano',
                          'Capital mediano*', 'Capital p80*', 'Capital p90*'])
order = ['A', 'B', 'C', 'D', 'E', 'D_snc', 'D_coord0', 'D_snc_coord0', 'D_collat100', 'D_cap11']
for i, k in enumerate(order):
    d = ST[k]; r_ = 5 + i
    put(ws, r_, 2, d['label'], BLACK, wrap=True)
    for j, (kk, fmt) in enumerate((('p_nogo', PCT), ('p_gate', PCT), ('p_fail', PCT), ('p_loss', PCT), ('p_loss_beyond_ph0', PCT), ('npv15_mean', NUM),
                                    ('moic_p50', MULT), ('cap_go_p50', NUM), ('cap_go_p80', NUM), ('cap_go_p90', NUM))):
        put(ws, r_, 3 + j, d[kk], BLACK, fmt)
r_ = 5 + len(order)
put(ws, r_, 2, '* Máximo acumulado que aportan los socios (aportes más fase 0, sin valor terminal), en los escenarios que pasan la fase 0.', NOTE)
r_ += 2
put(ws, r_, 2, 'Lo que arriesga un socio panameño (un sexto del capital y de las utilidades más un tercio de la fase 0; un tercio en la estructura E)', SUB); r_ += 1
hdr(ws, r_, range(2, 8), ['Estructura', 'Capital mediano si sigue', 'Capital p90 si sigue', 'Resultado neto promedio', 'Prob. de perder', 'Resultado neto p90'])
for k in order + ['D_cash70', 'D_cash80']:
    r_ += 1; d = ST[k]
    put(ws, r_, 2, d['label'], BLACK, wrap=True)
    for j, (kk, fmt) in enumerate((('j_cap_go_p50', NUM), ('j_cap_go_p90', NUM), ('j_net_mean', NUM), ('j_p_loss', PCT), ('j_net_p90', NUM))):
        put(ws, r_, 3 + j, d[kk], BLACK, fmt)
r_ += 2
put(ws, r_, 2, 'Supuestos de riesgo de la simulación', SUB); r_ += 1
for t in ['Planta con problemas de BPM: 15% de probabilidad; la fase 0 lo detecta 80% de las veces; si no, falla de calidad entre los meses 12 y 36, se ejecuta la fianza y se pierde el inventario.',
          'Ruptura del suministro por otras causas: 4% anual desde el mes 7, con ejecución de fianzas.',
          'Fase 0: se descarta el trato si se detecta el problema de planta o si el precio de la fábrica supera 1.6 veces la referencia (1.1 en la variante estricta).',
          'Compuerta operativa: 9 meses después de la primera venta, si las ventas anualizadas (más los contratos centralizados por arrancar) no llegan a B/.90,000, se cierra y no se hace la segunda ola.',
          'Correlaciones: participación descentralizada y probabilidad de ganar (0.6); mercado débil y mayor caída de precios (0.4); factor común de costos fijos (0.5).']:
    put(ws, r_, 2, t, BLACK, wrap=True); ws.merge_cells(start_row=r_, start_column=2, end_row=r_, end_column=11); ws.row_dimensions[r_].height = 28; r_ += 1

# =============================================================== Sensibilidad (fija)
ws = wb.create_sheet('Sensibilidad')
setw(ws, {'A': 4, 'B': 66, 'C': 14, 'D': 14, 'E': 14, 'F': 14, 'G': 14, 'H': 14, 'I': 14})
put(ws, 1, 2, 'Sensibilidad del VPN al 15% y del capital (un supuesto a la vez, el resto en la base)', TITLE)
put(ws, 2, 2, f"VPN base: B/.{A['tornado']['base_npv15']:,.0f}. Capital base: B/.{A['tornado']['base_cap']:,.0f}.", NOTE)
hdr(ws, 4, range(2, 10), ['Supuesto', 'Valor pesimista', 'Valor optimista', 'VPN pesimista', 'VPN optimista', 'Rango del VPN', 'Capital pesimista', 'Capital optimista'])
for i, t in enumerate(A['tornado']['rows']):
    r_ = 5 + i
    put(ws, r_, 2, t['label'], BLACK, wrap=True)
    put(ws, r_, 3, t['pes_val'], BLACK); put(ws, r_, 4, t['opt_val'], BLACK)
    put(ws, r_, 5, t['npv_pes'], BLACK, NUM); put(ws, r_, 6, t['npv_opt'], BLACK, NUM); put(ws, r_, 7, f'=F{r_}-E{r_}', BLACK, NUM)
    put(ws, r_, 8, t['cap_pes'], BLACK, NUM); put(ws, r_, 9, t['cap_opt'], BLACK, NUM)

# =============================================================== Vida útil (en vivo)
ws = wb.create_sheet('Vida útil')
setw(ws, {'A': 4, 'B': 60, 'C': 14})
put(ws, 1, 2, 'Presupuesto de vida útil (pliego de la LPU: vencimiento de la ficha técnica; 18 meses si no lo fija)', TITLE)
r_ = 3
put(ws, r_, 2, 'Vida útil aprobada del producto (meses)', BLACK); put(ws, r_, 3, 24, BLUE, '0', fill=YEL); sl_r = r_; r_ += 1
put(ws, r_, 2, 'Vida útil exigida al entregar (18 meses si la ficha técnica no fija otra)', BLACK); put(ws, r_, 3, 18, BLUE, '0'); mn_r = r_; r_ += 1
put(ws, r_, 2, 'Liberación en fábrica y despacho (meses)', BLACK); put(ws, r_, 3, 1, BLUE, '0.0'); a_r = r_; r_ += 1
put(ws, r_, 2, 'Tránsito y aduana (meses)', BLACK); put(ws, r_, 3, '=p_ship_m', GREEN, '0.0'); b_r = r_; r_ += 1
put(ws, r_, 2, 'Pedido a entrega en hospital (meses)', BLACK); put(ws, r_, 3, 0.25, BLUE, '0.00'); c_r = r_; r_ += 1
put(ws, r_, 2, 'Tiempo máximo que puede pasar en bodega (meses)', BOLD); cc = put(ws, r_, 3, f'=C{sl_r}-C{mn_r}-C{a_r}-C{b_r}-C{c_r}', BOLD, '0.00'); cc.fill = YEL; r_ += 2
put(ws, r_, 2, 'Con 24 meses de vida útil caben unos 2.75 meses en bodega: reposición mensual y poco inventario. Con 36 meses caben casi 15.', NOTE)

# =============================================================== Fuentes
ws = wb.create_sheet('Fuentes')
setw(ws, {'A': 4, 'B': 70, 'C': 100})
put(ws, 1, 2, 'Fuentes principales', TITLE)
src = [
 ('PanamaCompra: cuadro de propuestas de la LPU 01-2026 de la CSS', 'https://apisv3.panamacompra.gob.pa/documentos-actos-publico/cuadroPropuesta/12/procesoVistaCuadroPropuesta/712894'),
 ('Ley 419 de 2024 (medicamentos y compras públicas de medicamentos)', 'https://dnfd.minsa.gob.pa/sites/default/files/2025-08/Ley%20419%20de%201%20de%20febrero%20de%202024.pdf'),
 ('Ley 462 de 2025 (reforma de la CSS y de la Ley 419)', 'https://www.css.gob.pa/wp-content/uploads/2025/03/Ley-462-de-2025.pdf'),
 ('Ley 316 de 2022 (conflicto de intereses)', 'https://s3-legispan.asamblea.gob.pa/legispan/NORMAS/2020/2022/LEY/Administrador%20Legispan_29602-A_2022_8_18_ASAMBLEA%20NACIONAL_316.pdf'),
 ('Decreto Ejecutivo 27 de 2024 (tasas, licencias, registros, compras)', 'https://www.minsa.gob.pa/sites/default/files/general/decreto_ejecutivo_27_de_10_de_mayo_de_2024.pdf'),
 ('Resolución 985 de 2025 (principios activos con equivalencia terapéutica)', 'https://www.minsa.gob.pa/sites/default/files/general/resolucion_985_go30317_08072025.pdf'),
 ('PMBI: compras concluidas 2022-2025 (precios domésticos indios)', 'https://www.pmbi.co.in/Data/Awarded_Tender/Procurement%20Dept%202022_2025_13022026.pdf'),
 ('TNMSC: licitación e-018/2024 (precios domésticos indios)', 'https://tnmsc.tn.gov.in/linkfiles/news_events/N2502472.pdf'),
 ('Assam Medical Services Corporation: último precio de compra (2025)', 'https://amscl.assam.gov.in/sites/default/files/public_utility/Rates%20LPP-Drugs.pdf'),
 ('RBI: tasa de referencia INR/USD', 'https://www.rbi.org.in/'),
 ('AJOT: tarifas FAK de CMA CGM del subcontinente indio a Latinoamérica (jun-2026)', 'https://www.ajot.com/news/cma-cgm-fak-rates-from-the-indian-subcontinent-to-latin-america-and-north-europe-the-mediterranean-north-africa'),
 ('Banco General: tasas y comisiones de garantías', 'https://www.bgeneral.com/tasas-comisiones-recargos/'),
 ('Open Contracting: OCDS de la DGCP de Panamá (publicación 120)', 'https://data.open-contracting.org/en/publication/120'),
 ('Registro completo de verificaciones (474 afirmaciones con fuente y fecha)', 'research/verificacion_2026-10-09/registro_verificaciones.csv (repositorio)'),
]
hdr(ws, 3, range(2, 4), ['Fuente', 'Enlace'])
for i, (lab, url) in enumerate(src):
    put(ws, 4 + i, 2, lab, BLACK, wrap=True); put(ws, 4 + i, 3, url, BLACK)

wb.save(OUT)
print('saved', OUT, 'names', len(names), 'Mensual rows', R)
